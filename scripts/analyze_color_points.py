#!/usr/bin/env python3
"""Analyze red/green/blue point layouts in *_red/_green/_blue PNG files.

The script groups files by filename prefix, detects connected colored dots,
calculates counts/ratios/ranges/densities/clusters, and writes PNG/PDF/XLSX
outputs. It intentionally uses only Pillow, NumPy, OpenPyXL and ReportLab.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
from collections import Counter, defaultdict, deque
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

import numpy as np
from PIL import Image, ImageDraw, ImageFont


COLORS = ("red", "green", "blue")
RGB = {"red": (235, 35, 35), "green": (20, 190, 45), "blue": (35, 75, 220)}
SUFFIX_RE = re.compile(r"^(?P<sample>.+)_(?P<color>red|green|blue|combine)$", re.I)


@dataclass
class Point:
    sample: str
    color: str
    x: float
    y: float
    area: int
    diameter: float
    width: int
    height: int


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=Path(r"D:\Unet\half-section\output"))
    parser.add_argument("--output-dir", type=Path, default=None)
    parser.add_argument("--min-area", type=int, default=8, help="Minimum colored pixels in one dot")
    parser.add_argument("--max-area", type=int, default=2000, help="Maximum colored pixels in one dot")
    parser.add_argument("--threshold", type=int, default=120, help="Color channel threshold")
    parser.add_argument("--grid-size", type=int, default=80, help="Density grid width/height")
    parser.add_argument("--cluster-eps-px", type=float, default=45.0)
    parser.add_argument("--cluster-min-samples", type=int, default=8)
    parser.add_argument("--max-images", type=int, default=None, help="Debug: limit input image count")
    return parser.parse_args()


def sample_index(input_dir: Path, max_images: int | None) -> tuple[dict[str, dict[str, Path]], list[dict]]:
    groups: dict[str, dict[str, Path]] = defaultdict(dict)
    checks: list[dict] = []
    paths = sorted(input_dir.glob("*.png"))
    if max_images:
        paths = paths[:max_images]
    for path in paths:
        match = SUFFIX_RE.match(path.stem)
        if not match:
            checks.append({"file": path.name, "status": "unrecognized_name", "detail": "Expected *_red/green/blue/combine.png"})
            continue
        sample, color = match.group("sample"), match.group("color").lower()
        if color in groups[sample]:
            checks.append({"file": path.name, "status": "duplicate", "detail": f"Duplicate {color} for {sample}"})
        groups[sample][color] = path
    for sample, files in sorted(groups.items()):
        for color in COLORS:
            if color not in files:
                checks.append({"file": sample, "status": "missing", "detail": f"Missing {color}.png"})
        if "combine" not in files:
            checks.append({"file": sample, "status": "missing", "detail": "Missing combine.png"})
    return groups, checks


def color_mask(rgb: np.ndarray, color: str, threshold: int) -> np.ndarray:
    r, g, b = rgb[..., 0], rgb[..., 1], rgb[..., 2]
    if color == "red":
        return (r >= threshold) & (r > g * 1.35) & (r > b * 1.35)
    if color == "green":
        return (g >= threshold) & (g > r * 1.25) & (g > b * 1.15)
    return (b >= threshold) & (b > r * 1.35) & (b > g * 1.20)


def connected_components(mask: np.ndarray, min_area: int, max_area: int) -> list[tuple[int, int, int, float, float]]:
    """Return area, width, height, x-centre, y-centre for 8-connected components."""
    ys, xs = np.nonzero(mask)
    pixels = set(zip(xs.tolist(), ys.tolist()))
    components = []
    while pixels:
        start = pixels.pop()
        stack = [start]
        coords = [start]
        while stack:
            x, y = stack.pop()
            for nx in (x - 1, x, x + 1):
                for ny in (y - 1, y, y + 1):
                    if (nx, ny) in pixels:
                        pixels.remove((nx, ny))
                        stack.append((nx, ny))
                        coords.append((nx, ny))
        area = len(coords)
        if min_area <= area <= max_area:
            cx = sum(x for x, _ in coords) / area
            cy = sum(y for _, y in coords) / area
            min_x, max_x = min(x for x, _ in coords), max(x for x, _ in coords)
            min_y, max_y = min(y for _, y in coords), max(y for _, y in coords)
            components.append((area, max_x - min_x + 1, max_y - min_y + 1, cx, cy))
    return components


def extract_points(sample: str, color: str, path: Path, args: argparse.Namespace) -> tuple[list[Point], tuple[int, int] | None, str | None]:
    try:
        with Image.open(path) as image:
            rgb = np.asarray(image.convert("RGB"))
            width, height = image.size
    except Exception as exc:  # pragma: no cover - filesystem/image dependent
        return [], None, f"read_error: {exc}"
    return extract_points_from_rgb(sample, color, rgb, (width, height), args)


def extract_points_from_rgb(sample: str, color: str, rgb: np.ndarray, size: tuple[int, int], args: argparse.Namespace) -> tuple[list[Point], tuple[int, int] | None, str | None]:
    width, height = size
    mask = color_mask(rgb, color, args.threshold)
    components = connected_components(mask, args.min_area, args.max_area)
    points = [Point(sample, color, cx, cy, area, math.sqrt(4 * area / math.pi), w, h) for area, w, h, cx, cy in components]
    return points, (width, height), None


def dbscan(points: np.ndarray, eps: float, min_samples: int) -> np.ndarray:
    if len(points) == 0:
        return np.empty(0, dtype=int)
    # Spatial hash keeps DBSCAN practical for the full image collection.
    cell_size = max(float(eps), 1.0)
    buckets: dict[tuple[int, int], list[int]] = defaultdict(list)
    cells = np.floor(points / cell_size).astype(int)
    for i, cell in enumerate(cells):
        buckets[(int(cell[0]), int(cell[1]))].append(i)

    def neighbors(index: int) -> list[int]:
        cx, cy = int(cells[index, 0]), int(cells[index, 1])
        candidates: list[int] = []
        for dx in (-1, 0, 1):
            for dy in (-1, 0, 1):
                candidates.extend(buckets.get((cx + dx, cy + dy), []))
        if not candidates:
            return []
        candidate_array = np.asarray(candidates, dtype=int)
        distances = np.sqrt(((points[candidate_array] - points[index]) ** 2).sum(axis=1))
        return candidate_array[distances <= eps].tolist()

    labels = np.full(len(points), -2, dtype=int)  # -2 unvisited, -1 noise
    cluster_id = 0
    for i in range(len(points)):
        if labels[i] != -2:
            continue
        point_neighbors = neighbors(i)
        if len(point_neighbors) < min_samples:
            labels[i] = -1
            continue
        labels[i] = cluster_id
        queue = deque(point_neighbors)
        queued = set(point_neighbors)
        while queue:
            j = queue.popleft()
            if labels[j] == -1:
                labels[j] = cluster_id
            if labels[j] != -2:
                continue
            labels[j] = cluster_id
            expanded = neighbors(j)
            if len(expanded) >= min_samples:
                for k in expanded:
                    if k not in queued:
                        queue.append(k)
                        queued.add(k)
        cluster_id += 1
    return labels


def finite_stats(values: Iterable[float]) -> dict[str, float | None]:
    arr = np.asarray(list(values), dtype=float)
    if not len(arr):
        return {key: None for key in ("min", "max", "mean", "median", "std", "q25", "q75", "range")}
    return {"min": float(arr.min()), "max": float(arr.max()), "mean": float(arr.mean()), "median": float(np.median(arr)), "std": float(arr.std()), "q25": float(np.percentile(arr, 25)), "q75": float(np.percentile(arr, 75)), "range": float(arr.max() - arr.min())}


def make_canvas(width: int = 1200, height: int = 760) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (width, height), "white")
    return image, ImageDraw.Draw(image)


def font(size: int = 22):
    for candidate in ("C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/arial.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(candidate, size)
        except Exception:
            pass
    return ImageFont.load_default()


def transform_xy(x: float, y: float, image_size: tuple[int, int], canvas_size: tuple[int, int], margin: int = 70) -> tuple[float, float]:
    iw, ih = image_size
    cw, ch = canvas_size
    scale = min((cw - 2 * margin) / iw, (ch - 2 * margin) / ih)
    return margin + x * scale, margin + y * scale


def scatter_plot(points: list[Point], image_size: tuple[int, int], path: Path, title: str, clusters: np.ndarray | None = None) -> None:
    canvas, draw = make_canvas()
    draw.text((36, 22), title, fill="black", font=font(26))
    for i, point in enumerate(points):
        x, y = transform_xy(point.x, point.y, image_size, canvas.size)
        color = RGB[point.color]
        if clusters is not None and clusters[i] >= 0:
            color = tuple(min(255, int(v + int(255 * ((clusters[i] * 67) % 255) / 255) / 2)) for v in color)
        draw.ellipse((x - 4, y - 4, x + 4, y + 4), fill=color)
    canvas.save(path)


def density_plot(points: list[Point], image_size: tuple[int, int], path: Path, title: str, grid_size: int, color: str | None = None) -> None:
    canvas, draw = make_canvas()
    draw.text((36, 22), title, fill="black", font=font(26))
    selected = [p for p in points if color is None or p.color == color]
    hist = np.zeros((grid_size, grid_size), dtype=float)
    iw, ih = image_size
    for p in selected:
        ix = min(grid_size - 1, max(0, int(p.x / iw * grid_size)))
        iy = min(grid_size - 1, max(0, int(p.y / ih * grid_size)))
        hist[iy, ix] += 1
    max_count = float(hist.max()) if hist.size and hist.max() else 1.0
    margin = 70
    plot_w, plot_h = canvas.width - 2 * margin, canvas.height - 2 * margin
    base = RGB.get(color, (205, 40, 40)) if color else (150, 60, 180)
    for iy in range(grid_size):
        for ix in range(grid_size):
            intensity = float(hist[iy, ix] / max_count)
            if intensity <= 0:
                continue
            r = int(255 - (255 - base[0]) * intensity)
            g = int(255 - (255 - base[1]) * intensity)
            b = int(255 - (255 - base[2]) * intensity)
            x0 = margin + int(ix * plot_w / grid_size)
            y0 = margin + int(iy * plot_h / grid_size)
            x1 = margin + int((ix + 1) * plot_w / grid_size)
            y1 = margin + int((iy + 1) * plot_h / grid_size)
            draw.rectangle((x0, y0, x1, y1), fill=(r, g, b))
    canvas.save(path)


def bar_plot(counts: dict[str, int], path: Path, title: str) -> None:
    canvas, draw = make_canvas(1000, 700)
    draw.text((35, 25), title, fill="black", font=font(26))
    max_value = max(counts.values() or [1])
    base_y, top_y = 590, 100
    bar_width = 150
    for i, color in enumerate(COLORS):
        x0 = 160 + i * 260
        h = int((base_y - top_y) * counts.get(color, 0) / max_value)
        draw.rectangle((x0, base_y - h, x0 + bar_width, base_y), fill=RGB[color])
        draw.text((x0 + 35, base_y + 16), color, fill="black", font=font(22))
        draw.text((x0 + 48, base_y - h - 32), str(counts.get(color, 0)), fill="black", font=font(20))
    canvas.save(path)


def write_excel(output_dir: Path, points: list[Point], sample_rows: list[dict], distributions: list[dict], clusters: list[dict], checks: list[dict], params: dict) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    path = output_dir / "color_point_analysis.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Color point analysis", ""])
    ws.append(["Generated", datetime.now().isoformat(timespec="seconds")])
    ws.append(["Samples", len(sample_rows)])
    ws.append(["Detected points", len(points)])
    ws.append([])
    ws.append(["Color", "Points", "Share"])
    totals = Counter(p.color for p in points)
    total = max(1, len(points))
    for color in COLORS:
        ws.append([color, totals[color], totals[color] / total])

    def add_table(name: str, rows: list[dict]):
        sheet = wb.create_sheet(name)
        if not rows:
            sheet.append(["No records"])
            return
        headers = list(rows[0].keys())
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(h) for h in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4F81BD")
            cell.alignment = Alignment(horizontal="center")
        for col in range(1, sheet.max_column + 1):
            values = [str(sheet.cell(row, col).value or "") for row in range(1, min(sheet.max_row, 100) + 1)]
            sheet.column_dimensions[get_column_letter(col)].width = min(35, max(10, max(map(len, values), default=10) + 2))
        return sheet

    add_table("SampleSummary", sample_rows)
    point_rows = [{"sample": p.sample, "color": p.color, "x_px": p.x, "y_px": p.y, "x_norm": p.x / params["image_width"], "y_norm": p.y / params["image_height"], "area_px": p.area, "diameter_px": p.diameter, "width_px": p.width, "height_px": p.height} for p in points]
    add_table("PointCoordinates", point_rows)
    add_table("Distribution", distributions)
    add_table("Clusters", clusters)
    add_table("FileCheck", checks)
    add_table("Parameters", [{"parameter": k, "value": v} for k, v in params.items()])
    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
    wb.save(path)
    return path


def write_pdf(output_dir: Path, image_paths: list[Path], totals: dict[str, int], sample_count: int, point_count: int) -> Path | None:
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.pdfgen import canvas
    except Exception:
        return None
    path = output_dir / "color_point_analysis.pdf"
    page_w, page_h = landscape(A4)
    pdf = canvas.Canvas(str(path), pagesize=(page_w, page_h))
    pdf.setTitle("Color point analysis")
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(36, page_h - 40, "Color point analysis")
    pdf.setFont("Helvetica", 11)
    pdf.drawString(36, page_h - 62, f"Samples: {sample_count}    Points: {point_count}    " + "    ".join(f"{c}: {totals.get(c, 0)}" for c in COLORS))
    pdf.showPage()
    for image_path in image_paths:
        with Image.open(image_path) as image:
            iw, ih = image.size
        scale = min((page_w - 72) / iw, (page_h - 72) / ih)
        pdf.drawImage(str(image_path), 36, 36, width=iw * scale, height=ih * scale, preserveAspectRatio=True)
        pdf.showPage()
    pdf.save()
    return path


def main() -> int:
    args = parse_args()
    input_dir = args.input_dir.resolve()
    output_dir = (args.output_dir or input_dir.parent / "analysis_output").resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    if not input_dir.exists():
        print(f"Input directory does not exist: {input_dir}", file=sys.stderr)
        return 2

    groups, checks = sample_index(input_dir, args.max_images)
    points: list[Point] = []
    image_size: tuple[int, int] | None = None
    for sample, files in groups.items():
        if "combine" in files:
            try:
                with Image.open(files["combine"]) as image:
                    rgb = np.asarray(image.convert("RGB"))
                    size = image.size
                for color in COLORS:
                    if color in files:
                        found, _, _ = extract_points_from_rgb(sample, color, rgb, size, args)
                        points.extend(found)
                if image_size and size != image_size:
                    checks.append({"file": files["combine"].name, "status": "size_mismatch", "detail": f"{size} vs {image_size}"})
                image_size = image_size or size
            except Exception as exc:
                checks.append({"file": files["combine"].name, "status": "read_error", "detail": str(exc)})
        else:
            for color in COLORS:
                if color not in files:
                    continue
                found, size, error = extract_points(sample, color, files[color], args)
                points.extend(found)
                if size and image_size and size != image_size:
                    checks.append({"file": files[color].name, "status": "size_mismatch", "detail": f"{size} vs {image_size}"})
                image_size = image_size or size
                if error:
                    checks.append({"file": files[color].name, "status": "read_error", "detail": error})
    if image_size is None:
        print("No readable color PNG files found.", file=sys.stderr)
        return 3
    width, height = image_size

    sample_counts: dict[str, Counter] = defaultdict(Counter)
    for p in points:
        sample_counts[p.sample][p.color] += 1
    sample_rows = []
    for sample in sorted(groups):
        counts = sample_counts[sample]
        total = sum(counts[c] for c in COLORS)
        row = {"sample": sample, "sample_field_1": sample.split("_")[0], "sample_field_2": sample.split("_")[1] if len(sample.split("_")) > 1 else "", "sample_field_3": "_".join(sample.split("_")[2:]) if len(sample.split("_")) > 2 else "", "red": counts["red"], "green": counts["green"], "blue": counts["blue"], "total": total}
        for color in COLORS:
            row[f"{color}_share"] = counts[color] / total if total else 0
        sample_rows.append(row)

    distributions = []
    for color in COLORS:
        subset = [p for p in points if p.color == color]
        xs, ys = [p.x for p in subset], [p.y for p in subset]
        sx, sy = finite_stats(xs), finite_stats(ys)
        distributions.append({"color": color, "points": len(subset), "x_min_px": sx["min"], "x_max_px": sx["max"], "x_range_px": sx["range"], "x_mean_px": sx["mean"], "x_std_px": sx["std"], "y_min_px": sy["min"], "y_max_px": sy["max"], "y_range_px": sy["range"], "y_mean_px": sy["mean"], "y_std_px": sy["std"]})

    all_clusters: list[dict] = []
    cluster_plot_points = points
    if points:
        xy = np.array([[p.x, p.y] for p in points], dtype=float)
        labels = dbscan(xy, args.cluster_eps_px, args.cluster_min_samples)
        for idx, label in enumerate(labels):
            if label >= 0:
                all_clusters.append({"cluster_type": "mixed", "cluster_id": int(label), "sample": points[idx].sample, "color": points[idx].color, "x_px": points[idx].x, "y_px": points[idx].y, "noise": False})
            else:
                all_clusters.append({"cluster_type": "mixed", "cluster_id": -1, "sample": points[idx].sample, "color": points[idx].color, "x_px": points[idx].x, "y_px": points[idx].y, "noise": True})
        mixed_labels = labels
        for color in COLORS:
            indices = [i for i, p in enumerate(points) if p.color == color]
            if not indices:
                continue
            local_xy = xy[indices]
            local_labels = dbscan(local_xy, args.cluster_eps_px, args.cluster_min_samples)
            for local_i, point_i in enumerate(indices):
                p = points[point_i]
                label = int(local_labels[local_i])
                all_clusters.append({"cluster_type": color, "cluster_id": label if label >= 0 else -1, "sample": p.sample, "color": p.color, "x_px": p.x, "y_px": p.y, "noise": label < 0})
    else:
        mixed_labels = None

    totals = Counter(p.color for p in points)
    plot_paths: list[Path] = []
    all_path = output_dir / "01_all_points.png"
    scatter_plot(points, image_size, all_path, "All samples: red / green / blue points")
    plot_paths.append(all_path)
    for color in COLORS:
        path = output_dir / f"02_{color}_density.png"
        density_plot(points, image_size, path, f"{color.title()} point density", args.grid_size, color)
        plot_paths.append(path)
    count_path = output_dir / "03_color_counts.png"
    bar_plot(dict(totals), count_path, "Detected point counts")
    plot_paths.append(count_path)
    cluster_path = output_dir / "04_mixed_dbscan.png"
    scatter_plot(points, image_size, cluster_path, "Mixed-color DBSCAN clusters", mixed_labels)
    plot_paths.append(cluster_path)

    params = {"input_dir": str(input_dir), "output_dir": str(output_dir), "image_width": width, "image_height": height, "min_area": args.min_area, "max_area": args.max_area, "threshold": args.threshold, "grid_size": args.grid_size, "cluster_eps_px": args.cluster_eps_px, "cluster_min_samples": args.cluster_min_samples}
    xlsx_path = write_excel(output_dir, points, sample_rows, distributions, all_clusters, checks, params)
    pdf_path = write_pdf(output_dir, plot_paths, dict(totals), len(groups), len(points))
    report = {"generated": datetime.now().isoformat(timespec="seconds"), "samples": len(groups), "points": len(points), "counts": dict(totals), "image_size": image_size, "xlsx": str(xlsx_path), "pdf": str(pdf_path) if pdf_path else None, "plots": [str(p) for p in plot_paths], "checks": len(checks)}
    (output_dir / "run_summary.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
